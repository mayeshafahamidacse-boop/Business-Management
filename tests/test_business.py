import io
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from PIL import Image
from app import create_app
from app.db import get_db
from app.business import record_sale, add_payment, void_sale

class BusinessRules(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.data = Path(self.tmp.name) / 'data'
        self.app = create_app({'TESTING': True, 'DATA_DIR': str(self.data), 'WTF_CSRF_ENABLED': False})
        self.client = self.app.test_client()
        self.ctx = self.app.app_context(); self.ctx.push()
        d = get_db()
        d.executemany('insert into products(name,category,sku,cost,price,stock,threshold,created_at) values(?,?,?,?,?,?,?,?)', [
            ('Clip', 'Hair Clips', 'CLIP-1', 20, 50, 5, 2, '2026-10-01T10:00:00+06:00'),
            ('Scarf', 'Clothing', 'SCARF-1', 100, 200, 4, 1, '2026-10-01T10:00:00+06:00')])
        d.commit()
    def tearDown(self):
        self.ctx.pop(); self.tmp.cleanup()
    def test_multiline_snapshot_discount_stock_and_payment(self):
        d=get_db(); p=d.execute('select id from products order by id').fetchall()
        sid=record_sale(None,'Guest','','',[{'product_id':p[0][0],'quantity':2},{'product_id':p[1][0],'quantity':1}],20,100,'Cash')
        sale=d.execute('select * from sales where id=?',(sid,)).fetchone()
        self.assertEqual((sale['subtotal'],sale['discount'],sale['total']),(300,20,280))
        self.assertEqual(d.execute('select sum(quantity*unit_cost) from sale_items where sale_id=?',(sid,)).fetchone()[0],140)
        self.assertEqual([r[0] for r in d.execute('select stock from products order by id')],[3,3])
        d.execute('update products set cost=999,price=1500 where id=?',(p[0][0],));d.commit()
        self.assertEqual(tuple(d.execute('select unit_cost,unit_price from sale_items where sale_id=? order by id',(sid,)).fetchone()),(20,50))
        add_payment(sid,180,'bKash'); self.assertEqual(d.execute('select sum(amount) from payments where sale_id=?',(sid,)).fetchone()[0],280)
        with self.assertRaises(ValueError): add_payment(sid,1,'Cash')
    def test_oversell_rolls_back_and_void_restores_stock(self):
        d=get_db();p=d.execute('select id from products').fetchone()[0]
        with self.assertRaises(ValueError):record_sale(None,'','','',[{'product_id':p,'quantity':99}],0,0,'Cash')
        self.assertEqual(d.execute('select count(*) from sales').fetchone()[0],0)
        sid=record_sale(None,'','','',[{'product_id':p,'quantity':2}],0,0,'Cash')
        void_sale(sid);self.assertEqual(d.execute('select stock from products where id=?',(p,)).fetchone()[0],5)
        self.assertEqual(d.execute('select status from sales where id=?',(sid,)).fetchone()[0],'Voided')
    def test_product_search_and_photo_upload(self):
        image=io.BytesIO();Image.new('RGB',(12,12),'red').save(image,format='PNG');image.seek(0)
        response=self.client.post('/products',data={'name':'Ribbon','category':'Accessories','sku':'RIB-1','cost':'10','price':'25','stock':'8','threshold':'2','photo':(image,'ribbon.png')},content_type='multipart/form-data',follow_redirects=True)
        self.assertEqual(response.status_code,200)
        product=get_db().execute("select * from products where sku='RIB-1'").fetchone()
        self.assertTrue((self.data/'product_images'/product['image']).exists())
        image_response=self.client.get('/images/'+product['image']);self.assertEqual(image_response.status_code,200);image_response.close()
        self.assertIn(b'Ribbon',self.client.get('/products?q=Ribbon').data)
        response=self.client.post(f"/products/{product['id']}/edit",data={'name':'Ribbon XL','category':'Accessories','sku':'RIB-1','cost':'11','price':'28','stock':'7','threshold':'2','description':'Edited'})
        self.assertIn(response.status_code,(302,303))
        self.assertEqual(get_db().execute('select name,price from products where id=?',(product['id'],)).fetchone()[:],('Ribbon XL',28))
    def test_report_and_real_excel_export(self):
        d=get_db();p=d.execute('select id from products').fetchone()[0]
        record_sale(None,'Guest','','',[{'product_id':p,'quantity':1}],0,10,'Cash')
        self.assertEqual(self.client.get('/reports?month=2026-10').status_code,200)
        response=self.client.get('/export/monthly?month=2026-10')
        wb=load_workbook(io.BytesIO(response.data),read_only=True)
        self.assertIn('Summary',wb.sheetnames);self.assertIn('Sales',wb.sheetnames);self.assertIn('Category sales',wb.sheetnames);self.assertIn('Payment methods',wb.sheetnames)
    def test_backup_and_restore(self):
        self.client.post('/backup')
        backup=next((self.data/'backups').glob('*.zip'))
        get_db().execute("insert into customers(name,created_at) values('Temporary','2026-10-21')");get_db().commit()
        from flask import g
        old=g.pop('db',None)
        if old: old.close()
        with backup.open('rb') as stream:
            response=self.client.post('/restore',data={'backupfile':(stream,'backup.zip')},content_type='multipart/form-data')
        self.assertIn(response.status_code,(302,303))
        self.assertEqual(get_db().execute("select count(*) from customers where name='Temporary'").fetchone()[0],0)
        self.assertTrue(list((self.data/'backups').glob('before_restore_*.zip')))

if __name__=='__main__': unittest.main()




