from flask import Blueprint,render_template,request,redirect,url_for,flash
from datetime import datetime,date
from .db import get_db
from .business import now,record_sale,local_today
bp=Blueprint('main',__name__)
@bp.app_context_processor
def common_template_values(): return {'categories':['Clothing','Hair Clips','Accessories','Jewelry','Ornaments','Other']}
@bp.route('/')
def dashboard():
 from datetime import date
 d=get_db();start=request.args.get('from',local_today().replace(day=1).isoformat());finish=request.args.get('to',local_today().isoformat())
 sales=d.execute("select coalesce(sum(total),0) rev,count(*) orders from sales where status='Active' and date(created_at) between ? and ?",(start,finish)).fetchone()
 paid=d.execute("select coalesce(sum(amount),0) from payments where date(paid_at) between ? and ?",(start,finish)).fetchone()[0]
 cogs=d.execute("select coalesce(sum(i.quantity*i.unit_cost),0),coalesce(sum(i.quantity),0) from sale_items i join sales s on s.id=i.sale_id where s.status='Active' and date(s.created_at) between ? and ?",(start,finish)).fetchone()
 bal=d.execute("select coalesce(sum(s.total-(select coalesce(sum(amount),0) from payments where sale_id=s.id)),0) from sales s where status='Active'").fetchone()[0]
 low=d.execute('select count(*) from products where stock<=threshold').fetchone()[0]
 best=d.execute("select i.product_name,sum(i.quantity) qty from sale_items i join sales s on s.id=i.sale_id where s.status='Active' and date(s.created_at) between ? and ? group by i.product_name order by qty desc limit 5",(start,finish)).fetchall()
 today_total=d.execute("select coalesce(sum(total),0) from sales where status='Active' and date(created_at)=?",(local_today().isoformat(),)).fetchone()[0]
 return render_template('dashboard.html',title='Dashboard',begin=start,finish=finish,today_total=today_total,revenue=sales['rev'],orders=sales['orders'],paid=paid,cogs=cogs[0],units=cogs[1],profit=sales['rev']-cogs[0],receivable=bal,low=low,best=best)
@bp.route('/products',methods=['GET','POST'])
def products():
 from PIL import Image,UnidentifiedImageError
 from werkzeug.utils import secure_filename
 from pathlib import Path
 from flask import current_app
 import sqlite3
 d=get_db()
 if request.method=='POST':
  image=None;f=request.files.get('photo')
  try:
   if f and f.filename:
    im=Image.open(f.stream);im.verify();f.stream.seek(0);im=Image.open(f.stream);im.thumbnail((1000,1000));image=secure_filename(f.filename.rsplit('.',1)[0])+'.jpg';im.convert('RGB').save(Path(current_app.config['DATA_DIR'])/'product_images'/image,'JPEG',quality=88)
   values=(request.form['name'].strip(),request.form['category'],request.form.get('sku') or None,image,float(request.form['cost']),float(request.form['price']),int(request.form['stock']),int(request.form.get('threshold',5)),request.form.get('description',''),now())
   if min(values[4],values[5],values[6],values[7])<0:raise ValueError('Amounts and stock cannot be negative')
   d.execute('insert into products(name,category,sku,image,cost,price,stock,threshold,description,created_at) values(?,?,?,?,?,?,?,?,?,?)',values);d.commit();flash('Product saved')
  except (ValueError,sqlite3.Error,UnidentifiedImageError,OSError) as e:d.rollback();flash(str(e),'error')
  return redirect(url_for('main.products'))
 q=request.args.get('q','');cat=request.args.get('category','')
 products=d.execute("select * from products where (name like ? or coalesce(sku,'') like ?) and (?='' or category=?) order by name",('%'+q+'%','%'+q+'%',cat,cat)).fetchall()
 return render_template('products.html',title='Products',products=products,q=q,cat=cat)
@bp.route('/customers',methods=['GET','POST'])
def customers():
 d=get_db()
 if request.method=='POST':
  try:d.execute('insert into customers(name,phone,address,notes,created_at) values(?,?,?,?,?)',(request.form['name'].strip(),request.form.get('phone'),request.form.get('address'),request.form.get('notes'),now()));d.commit();flash('Customer saved')
  except Exception as e:d.rollback();flash(str(e),'error')
 q=request.args.get('q','')
 rows=d.execute("select c.*,coalesce((select sum(s.total) from sales s where s.customer_id=c.id and s.status='Active'),0) spent,coalesce((select sum(s.total-(select coalesce(sum(p.amount),0) from payments p where p.sale_id=s.id)) from sales s where s.customer_id=c.id and s.status='Active'),0) balance from customers c where c.name like ? or coalesce(c.phone,'') like ? order by c.name",('%'+q+'%','%'+q+'%')).fetchall()
 return render_template('customers.html',title='Customers',customers=rows,q=q)
@bp.route('/sales/new',methods=['GET','POST'])
def new_sale():
 d=get_db()
 if request.method=='POST':
  try:
   ids=request.form.getlist('product_id[]');qs=request.form.getlist('quantity[]');prices=request.form.getlist('unit_price[]');rows=[{'product_id':i,'quantity':q,'unit_price':price} for i,q,price in zip(ids,qs,prices) if i and q]
   sid=record_sale(request.form.get('customer_id') or None,request.form.get('guest_name',''),request.form.get('phone',''),request.form.get('address',''),rows,request.form.get('discount',0),request.form.get('paid',0),request.form.get('method','Cash'),request.form.get('notes',''));return redirect(url_for('main.sale_detail',sid=sid))
  except Exception as e:flash(str(e),'error')
 return render_template('new_sale.html',title='New Sale',products=d.execute('select * from products where stock>0').fetchall(),customers=d.execute('select * from customers').fetchall())
@bp.route('/sales')
def sales():
 d=get_db();q=request.args.get('q','');start=request.args.get('from','');end=request.args.get('to','')
 sql="select s.*,coalesce((select sum(amount) from payments where sale_id=s.id),0) paid,(select group_concat(product_name||' x'||quantity,', ') from sale_items where sale_id=s.id) items from sales s where (s.customer_name like ? or coalesce(s.customer_phone,'') like ? or cast(s.id as text) like ?) and (?='' or date(s.created_at)>=?) and (?='' or date(s.created_at)<=?) order by s.id desc"
 rows=d.execute(sql,('%'+q+'%','%'+q+'%','%'+q+'%',start,start,end,end)).fetchall();return render_template('sales.html',title='Sales History',sales=rows,q=q,start=start,end=end)
@bp.route('/reports')
def reports():
 d=get_db();m=request.args.get('month',local_today().strftime('%Y-%m'))
 try:a=datetime.strptime(m+'-01','%Y-%m-%d');b=datetime(a.year+a.month//12,a.month%12+1,1)
 except ValueError:m=local_today().strftime('%Y-%m');a=datetime.strptime(m+'-01','%Y-%m-%d');b=datetime(a.year+a.month//12,a.month%12+1,1)
 bounds=(a.isoformat(),b.isoformat())
 totals=d.execute("select coalesce(sum(subtotal),0) gross,coalesce(sum(discount),0) discount,coalesce(sum(total),0) net,count(*) orders from sales where status='Active' and created_at>=? and created_at<?",bounds).fetchone()
 cogs=d.execute("select coalesce(sum(i.quantity*i.unit_cost),0),coalesce(sum(i.quantity),0) from sale_items i join sales s on s.id=i.sale_id where s.status='Active' and s.created_at>=? and s.created_at<?",bounds).fetchone()
 received=d.execute("select coalesce(sum(amount),0) from payments where paid_at>=? and paid_at<?",bounds).fetchone()[0]
 outstanding=d.execute("select coalesce(sum(s.total-(select coalesce(sum(amount),0) from payments where sale_id=s.id)),0) from sales s where s.status='Active' and s.created_at>=? and s.created_at<?",bounds).fetchone()[0]
 daily=d.execute("select date(created_at) day,sum(total) total from sales where status='Active' and created_at>=? and created_at<? group by date(created_at) order by day",bounds).fetchall()
 category=d.execute("select i.category,sum(i.quantity*i.unit_price) revenue from sale_items i join sales s on s.id=i.sale_id where s.status='Active' and s.created_at>=? and s.created_at<? group by i.category order by revenue desc",bounds).fetchall()
 methods=d.execute("select method,sum(amount) received from payments where paid_at>=? and paid_at<? group by method order by received desc",bounds).fetchall()
 best=d.execute("select i.product_name,sum(i.quantity) units from sale_items i join sales s on s.id=i.sale_id where s.status='Active' and s.created_at>=? and s.created_at<? group by i.product_name order by units desc limit 10",bounds).fetchall()
 return render_template('reports.html',title='Monthly Reports',month=m,reports=[totals],cogs=cogs[0],units=cogs[1],received=received,outstanding=outstanding,daily=daily,category=category,methods=methods,best=best,average=(totals['net']/totals['orders'] if totals['orders'] else 0))
@bp.route('/exports')
def exports():return render_template('exports.html',title='Excel Exports')
@bp.route('/inventory')
def inventory():return redirect(url_for('main.products'))
@bp.route('/settings')
def settings():return render_template('settings.html',title='Settings')
@bp.route('/backup',methods=['GET','POST'])
def backup():
 import zipfile
 from pathlib import Path
 from flask import current_app
 data=Path(current_app.config['DATA_DIR']); folder=data/'backups'; folder.mkdir(exist_ok=True)
 if request.method=='POST':
  target=folder/('boutique_backup_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.zip')
  with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
   for f in data.rglob('*'):
    if f.is_file() and 'backups' not in f.parts:z.write(f,f.relative_to(data))
  flash('Backup created: '+str(target))
 return render_template('backup.html',title='Backup & Restore')
@bp.route('/sales/<int:sid>',methods=['GET','POST'])
def sale_detail(sid):
 from .business import add_payment
 d=get_db()
 if request.method=='POST':
  try:add_payment(sid,request.form['amount'],request.form['method']);flash('Payment recorded')
  except Exception as e:flash(str(e),'error')
  return redirect(url_for('main.sale_detail',sid=sid))
 s=d.execute('select * from sales where id=?',(sid,)).fetchone()
 if not s:return 'Not found',404
 return render_template('sale_detail.html',title='Sale details',sale=s,items=d.execute('select * from sale_items where sale_id=?',(sid,)).fetchall(),payments=d.execute('select * from payments where sale_id=?',(sid,)).fetchall())
@bp.route('/sales/<int:sid>/void',methods=['POST'])
def void_order(sid):
 from .business import void_sale
 try:void_sale(sid);flash('Sale voided and stock restored')
 except Exception as e:flash(str(e),'error')
 return redirect(url_for('main.sales'))
@bp.route('/export/<kind>')
def export(kind):
 import io
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 from openpyxl.utils import get_column_letter
 from flask import send_file
 choices={'sales':('select id,created_at,customer_name,customer_phone,subtotal,discount,total,method,status from sales',['ID','Date','Customer','Phone','Gross','Discount','Net','Method','Status']),'items':('select s.id,s.created_at,s.customer_name,i.product_name,i.quantity,i.unit_price,i.unit_cost from sale_items i join sales s on s.id=i.sale_id',['Sale ID','Date','Customer','Product','Qty','Price','Cost']),'inventory':('select id,name,category,sku,cost,price,stock,threshold from products',['ID','Name','Category','SKU','Cost','Price','Stock','Threshold']),'customers':('select id,name,phone,address,notes from customers',['ID','Name','Phone','Address','Notes'])}
 if kind not in choices:return 'Not found',404
 sql,heads=choices[kind]
 start=request.args.get('from','');end=request.args.get('to','');params=()
 if kind in ('sales','items'):
  date_col='created_at' if kind=='sales' else 's.created_at'
  sql+=f" where (? is null or date({date_col})>=date(?)) and (? is null or date({date_col})<=date(?))"
  params=(start or None,start or None,end or None,end or None)
 w=Workbook();s=w.active;s.append(heads)
 for row in get_db().execute(sql,params):s.append(list(row))
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='385A52')
 for col in s.columns:s.column_dimensions[get_column_letter(col[0].column)].width=18
 f=io.BytesIO();w.save(f);f.seek(0);return send_file(f,as_attachment=True,download_name='boutique_'+kind+'.xlsx')

import sqlite3,zipfile,shutil
from pathlib import Path
from flask import current_app,send_file
@bp.route('/restore',methods=['POST'])
def restore_backup():
 data=Path(current_app.config['DATA_DIR']); upload=request.files.get('backupfile')
 if not upload:flash('Select a backup ZIP.','error');return redirect(url_for('main.backup'))
 stage=data/'restore_staging';shutil.rmtree(stage,ignore_errors=True);stage.mkdir()
 try:
  with zipfile.ZipFile(upload) as z:
   if any(Path(n).is_absolute() or '..' in Path(n).parts for n in z.namelist()):raise ValueError('Unsafe backup paths')
   z.extractall(stage)
  candidate=stage/'business.sqlite3'
  if not candidate.exists():raise ValueError('Backup database is missing')
  c=sqlite3.connect(candidate);check=c.execute('pragma integrity_check').fetchone()[0];c.close()
  if check!='ok':raise ValueError('Database integrity check failed')
  # Preserve the current dataset before replacing any files.
  before=data/'backups';before.mkdir(exist_ok=True);snapshot=before/('before_restore_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.zip')
  with zipfile.ZipFile(snapshot,'w',zipfile.ZIP_DEFLATED) as z:
   for f in data.rglob('*'):
    if f.is_file() and 'backups' not in f.parts and 'restore_staging' not in f.parts:z.write(f,f.relative_to(data))
  # Close this request's handle before replacing the database; restart the server after restoring.
  from flask import g
  old=g.pop('db',None)
  if old:old.close()
  for child in stage.iterdir():
   dest=data/child.name
   if dest.is_dir():shutil.rmtree(dest)
   elif dest.exists():dest.unlink()
   shutil.move(str(child),str(dest))
  flash('Restore completed. Restart the app to load the restored records.')
 except (zipfile.BadZipFile,sqlite3.Error,ValueError,OSError) as exc:flash('Restore failed: '+str(exc),'error')
 finally:shutil.rmtree(stage,ignore_errors=True)
 return redirect(url_for('main.backup'))


@bp.route('/export/monthly')
def export_monthly():
 import io
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 from flask import send_file
 m=request.args.get('month',local_today().strftime('%Y-%m'))
 try: start=datetime.strptime(m+'-01','%Y-%m-%d');end=datetime(start.year+start.month//12,start.month%12+1,1)
 except ValueError:return 'Invalid month',400
 d=get_db(); a=start.isoformat();b=end.isoformat()
 totals=d.execute("select coalesce(sum(subtotal),0),coalesce(sum(discount),0),coalesce(sum(total),0),count(*) from sales where status='Active' and created_at>=? and created_at<?",(a,b)).fetchone()
 cogs=d.execute("select coalesce(sum(i.quantity*i.unit_cost),0),coalesce(sum(i.quantity),0) from sale_items i join sales s on s.id=i.sale_id where s.status='Active' and s.created_at>=? and s.created_at<?",(a,b)).fetchone()
 received=d.execute("select coalesce(sum(amount),0) from payments where paid_at>=? and paid_at<?",(a,b)).fetchone()[0]
 outstanding=d.execute("select coalesce(sum(s.total-(select coalesce(sum(amount),0) from payments where sale_id=s.id)),0) from sales s where s.status='Active' and s.created_at>=? and s.created_at<?",(a,b)).fetchone()[0]
 w=Workbook();summary=w.active;summary.title='Summary';summary.append(['Monthly report',m]);summary.append(['Gross sales',totals[0]]);summary.append(['Discounts',totals[1]]);summary.append(['Net sales',totals[2]]);summary.append(['COGS',cogs[0]]);summary.append(['Gross profit',totals[2]-cogs[0]]);summary.append(['Payments received',received]);summary.append(['Outstanding for month',outstanding]);summary.append(['Orders',totals[3]]);summary.append(['Units sold',cogs[1]])
 detail=w.create_sheet('Sales');detail.append(['ID','Date','Customer','Gross','Discount','Net','Received','Balance','Status'])
 for r in d.execute("select s.id,s.created_at,s.customer_name,s.subtotal,s.discount,s.total,coalesce((select sum(amount) from payments where sale_id=s.id),0),s.status from sales s where s.created_at>=? and s.created_at<? order by s.id",(a,b)):
  detail.append([*r[:6],r[6],r[5]-r[6],r[7]])
 categories=w.create_sheet('Category sales');categories.append(['Category','Net item sales'])
 for row in d.execute("select i.category,sum(i.quantity*i.unit_price) from sale_items i join sales s on s.id=i.sale_id where s.status='Active' and s.created_at>=? and s.created_at<? group by i.category",(a,b)): categories.append(list(row))
 methods=w.create_sheet('Payment methods');methods.append(['Method','Amount received'])
 for row in d.execute("select p.method,sum(p.amount) from payments p join sales s on s.id=p.sale_id where s.status='Active' and p.paid_at>=? and p.paid_at<? group by p.method",(a,b)): methods.append(list(row))
 for sh in w:
  sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
  for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='385A52')
  for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=22
 out=io.BytesIO();w.save(out);out.seek(0);return send_file(out,as_attachment=True,download_name='boutique_month_'+m+'.xlsx')

@bp.route('/products/<int:pid>/edit',methods=['GET','POST'])
def edit_product(pid):
 from PIL import Image,UnidentifiedImageError
 from werkzeug.utils import secure_filename
 from pathlib import Path
 from flask import current_app,send_from_directory
 import sqlite3
 d=get_db();p=d.execute('select * from products where id=?',(pid,)).fetchone()
 if not p:return 'Not found',404
 if request.method=='POST':
  try:
   image=p['image'];f=request.files.get('photo')
   if f and f.filename:
    im=Image.open(f.stream);im.verify();f.stream.seek(0);im=Image.open(f.stream);im.thumbnail((1000,1000));image=secure_filename(f.filename.rsplit('.',1)[0])+'.jpg';im.convert('RGB').save(Path(current_app.config['DATA_DIR'])/'product_images'/image,'JPEG',quality=88)
   vals=(request.form['name'],request.form['category'],request.form.get('sku') or None,float(request.form['cost']),float(request.form['price']),int(request.form['stock']),int(request.form['threshold']),request.form.get('description',''),image,pid)
   if min(vals[3],vals[4],vals[5],vals[6])<0:raise ValueError('Values cannot be negative')
   d.execute('update products set name=?,category=?,sku=?,cost=?,price=?,stock=?,threshold=?,description=?,image=? where id=?',vals);d.commit();flash('Product updated');return redirect(url_for('main.products'))
  except (ValueError,sqlite3.Error,UnidentifiedImageError,OSError) as e:d.rollback();flash(str(e),'error')
 return render_template('edit_product.html',title='Edit product',product=p)
@bp.route('/products/<int:pid>/delete',methods=['POST'])
def delete_product(pid):
 d=get_db()
 try:d.execute('delete from products where id=?',(pid,));d.commit();flash('Product removed; historical sale item snapshots remain.')
 except Exception as e:d.rollback();flash(str(e),'error')
 return redirect(url_for('main.products'))
@bp.route('/customers/<int:cid>/edit',methods=['GET','POST'])
def edit_customer(cid):
 d=get_db();c=d.execute('select * from customers where id=?',(cid,)).fetchone()
 if not c:return 'Not found',404
 if request.method=='POST':
  d.execute('update customers set name=?,phone=?,address=?,notes=? where id=?',(request.form['name'],request.form.get('phone'),request.form.get('address'),request.form.get('notes'),cid));d.commit();flash('Customer updated');return redirect(url_for('main.customers'))
 return render_template('edit_customer.html',title='Edit customer',customer=c)
@bp.route('/customers/<int:cid>')
def customer_detail(cid):
 d=get_db();c=d.execute('select * from customers where id=?',(cid,)).fetchone()
 if not c:return 'Not found',404
 sales=d.execute("select s.*,coalesce((select sum(amount) from payments where sale_id=s.id),0) paid from sales s where customer_id=? order by id desc",(cid,)).fetchall()
 return render_template('customer_detail.html',title='Customer history',customer=c,sales=sales)

from datetime import date




@bp.route('/images/<path:name>')
def product_image(name):
 from pathlib import Path
 from flask import current_app,send_from_directory
 return send_from_directory(Path(current_app.config['DATA_DIR'])/'product_images',name)

@bp.route('/export/balances')
def export_balances():
 import io
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 from flask import send_file
 d=get_db();w=Workbook();s=w.active;s.title='Outstanding';s.append(['Sale ID','Date','Customer','Phone','Sale total','Payments received','Outstanding'])
 for r in d.execute("select s.id,s.created_at,s.customer_name,s.customer_phone,s.total,coalesce((select sum(amount) from payments where sale_id=s.id),0) from sales s where s.status='Active' order by s.created_at"):
  s.append([r[0],r[1],r[2],r[3],r[4],r[5],r[4]-r[5]])
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='385A52')
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=22
 out=io.BytesIO();w.save(out);out.seek(0);return send_file(out,as_attachment=True,download_name='boutique_outstanding.xlsx')





@bp.route('/sales/<int:sid>/edit',methods=['POST'])
def edit_sale_info(sid):
 d=get_db();sale=d.execute("select status from sales where id=?",(sid,)).fetchone()
 if not sale or sale['status']!='Active':flash('Only active sales can be edited.','error');return redirect(url_for('main.sales'))
 name=request.form.get('customer_name','').strip() or 'Walk-in'
 d.execute('update sales set customer_name=?,customer_phone=?,customer_address=?,notes=? where id=?',(name,request.form.get('customer_phone','').strip(),request.form.get('customer_address','').strip(),request.form.get('notes','').strip(),sid));d.commit();flash('Sale and customer contact details updated.');return redirect(url_for('main.sale_detail',sid=sid))

